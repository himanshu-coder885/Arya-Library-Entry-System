import csv
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from library_app.config import (
    BASE_DIR,
    DEFAULT_STUDENTS_FILE,
    EXCEL_STUDENTS_FILE,
    LIBRARY_DATA_FILE,
    LIBRARY_DB_FILE,
    VISITS_FILE,
)


def get_connection():
    conn = sqlite3.connect(LIBRARY_DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn


def initialize_database():
    with get_connection() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS students (
                student_id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                father_name TEXT,
                course TEXT,
                phone TEXT,
                valid_until TEXT
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS visits (
                visit_id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id TEXT NOT NULL,
                name TEXT NOT NULL,
                father_name TEXT,
                date TEXT NOT NULL,
                entry_time TEXT NOT NULL,
                exit_time TEXT DEFAULT '',
                FOREIGN KEY(student_id) REFERENCES students(student_id)
            )
            """
        )
        student_columns = {row["name"] for row in conn.execute("PRAGMA table_info(students)").fetchall()}
        if "father_name" not in student_columns:
            conn.execute("ALTER TABLE students ADD COLUMN father_name TEXT DEFAULT ''")

        visit_columns = {row["name"] for row in conn.execute("PRAGMA table_info(visits)").fetchall()}
        if "father_name" not in visit_columns:
            conn.execute("ALTER TABLE visits ADD COLUMN father_name TEXT DEFAULT ''")


def _normalize_student_row(row):
    return {
        "student_id": (row.get("student_id") or row.get("Student ID") or "").strip(),
        "name": (row.get("name") or row.get("Name") or "").strip(),
        "father_name": (row.get("father_name") or row.get("Father Name") or row.get("FATHER NAME") or "").strip(),
        "course": (
            row.get("course")
            or row.get("Course")
            or row.get("coursev1")
            or row.get("Coursev1")
            or row.get("branch")
            or row.get("Branch")
            or ""
        ).strip(),
        "phone": (row.get("phone") or row.get("Phone") or "").strip(),
        "valid_until": (row.get("valid_until") or row.get("Valid Until") or "").strip(),
    }


def _student_source_file():
    if EXCEL_STUDENTS_FILE.exists():
        return EXCEL_STUDENTS_FILE
    if LIBRARY_DATA_FILE.exists():
        return LIBRARY_DATA_FILE
    return DEFAULT_STUDENTS_FILE


def _excel_source_files():
    preferred = []
    if EXCEL_STUDENTS_FILE.exists():
        preferred.append(EXCEL_STUDENTS_FILE)

    others = sorted(
        [
            path
            for path in BASE_DIR.glob("*.xlsx")
            if path.name != EXCEL_STUDENTS_FILE.name
        ]
    )
    return preferred + others


def import_students_from_excel():
    excel_files = _excel_source_files()
    if not excel_files:
        return False

    students = {}

    for excel_file in excel_files:
        workbook = load_workbook(excel_file, read_only=True, data_only=True)

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)

            headers = None
            for row in rows:
                values = [str(cell).strip() if cell is not None else "" for cell in row]
                if not any(values):
                    continue

                normalized = [value.lower() for value in values]
                if headers is None:
                    if (
                        (
                            "name" in normalized
                            or "name " in normalized
                            or "studente name" in normalized
                            or "students name" in normalized
                        )
                        and "branch" in normalized
                        and "code" in normalized
                    ):
                        headers = values
                    continue

                if headers is None:
                    continue

                row_map = dict(zip(headers, values))
                student_id = str(row_map.get("CODE", "")).strip()
                name = str(
                    row_map.get("Name ", "")
                    or row_map.get("Name", "")
                    or row_map.get("STUDENTE NAME", "")
                    or row_map.get("Students Name", "")
                ).strip()
                course = str(row_map.get("BRANCH", "") or row_map.get("Branch", "")).strip()

                if not student_id or not name:
                    continue

                students[student_id] = {
                    "student_id": student_id,
                    "name": name,
                    "father_name": str(
                        row_map.get("FATHER NAME", "")
                        or row_map.get("Father Name", "")
                    ).strip(),
                    "course": course,
                    "phone": "",
                    "valid_until": "",
                }

    with get_connection() as conn:
        conn.execute("DELETE FROM students")
        for student in students.values():
            conn.execute(
                """
                INSERT INTO students(student_id, name, father_name, course, phone, valid_until)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    student["student_id"],
                    student["name"],
                    student["father_name"],
                    student["course"],
                    student["phone"],
                    student["valid_until"],
                ),
            )

    return True


def import_students_from_csv():
    source = _student_source_file()
    if source.suffix.lower() == ".xlsx":
        return
    if not source.exists():
        return

    with source.open("r", newline="", encoding="utf-8") as file, get_connection() as conn:
        reader = csv.DictReader(file)
        for row in reader:
            student = _normalize_student_row(row)
            if not student["student_id"]:
                continue
            conn.execute(
                """
                INSERT INTO students(student_id, name, father_name, course, phone, valid_until)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(student_id) DO UPDATE SET
                    name=excluded.name,
                    father_name=excluded.father_name,
                    course=excluded.course,
                    phone=excluded.phone,
                    valid_until=excluded.valid_until
                """,
                (
                    student["student_id"],
                    student["name"],
                    student["father_name"],
                    student["course"],
                    student["phone"],
                    student["valid_until"],
                ),
            )


def import_visits_from_csv():
    if not VISITS_FILE.exists():
        return

    with VISITS_FILE.open("r", newline="", encoding="utf-8") as file, get_connection() as conn:
        count = conn.execute("SELECT COUNT(*) FROM visits").fetchone()[0]
        if count:
            return

        for row in csv.DictReader(file):
            student_id = (row.get("student_id") or "").strip()
            name = (row.get("name") or "").strip()
            date = (row.get("date") or row.get("scan_date") or "").strip()
            entry_time = (row.get("entry_time") or "").strip()
            exit_time = (row.get("exit_time") or "").strip()
            father_name = (row.get("father_name") or "").strip()
            if not student_id or not date or not entry_time:
                continue
            conn.execute(
                """
                INSERT INTO visits(student_id, name, father_name, date, entry_time, exit_time)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (student_id, name, father_name, date, entry_time, exit_time),
            )


def ensure_database_ready():
    initialize_database()
    if not import_students_from_excel():
        import_students_from_csv()
    import_visits_from_csv()


def fetch_students():
    ensure_database_ready()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT student_id, name, father_name, course, phone, valid_until FROM students ORDER BY student_id"
        ).fetchall()
    return {
        row["student_id"]: {
            "student_id": row["student_id"],
            "name": row["name"],
            "father_name": row["father_name"] or "",
            "course": row["course"] or "",
            "phone": row["phone"] or "",
            "valid_until": row["valid_until"] or "",
        }
        for row in rows
    }


def fetch_visits():
    ensure_database_ready()
    with get_connection() as conn:
        rows = conn.execute(
            """
            SELECT visit_id, student_id, name, father_name, date, entry_time, exit_time
            FROM visits
            ORDER BY visit_id
            """
        ).fetchall()
    return [
        {
            "visit_id": str(row["visit_id"]).zfill(5),
            "student_id": row["student_id"],
            "name": row["name"],
            "father_name": row["father_name"] or "",
            "date": row["date"],
            "entry_time": row["entry_time"],
            "exit_time": row["exit_time"] or "",
        }
        for row in rows
    ]


def create_visit(student):
    now = datetime.now()
    ensure_database_ready()
    with get_connection() as conn:
        cursor = conn.execute(
            """
            INSERT INTO visits(student_id, name, father_name, date, entry_time, exit_time)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                student["student_id"],
                student["name"],
                student.get("father_name", ""),
                now.strftime("%Y-%m-%d"),
                now.strftime("%H:%M:%S"),
                "",
            ),
        )
        visit_id = cursor.lastrowid

    return {
        "visit_id": str(visit_id).zfill(5),
        "student_id": student["student_id"],
        "name": student["name"],
        "father_name": student.get("father_name", ""),
        "date": now.strftime("%Y-%m-%d"),
        "entry_time": now.strftime("%H:%M:%S"),
        "exit_time": "",
    }


def update_visit_exit(student_id, visit_date=None):
    now = datetime.now()
    target_date = visit_date or now.strftime("%Y-%m-%d")
    ensure_database_ready()
    with get_connection() as conn:
        row = conn.execute(
            """
            SELECT visit_id, student_id, name, father_name, date, entry_time, exit_time
            FROM visits
            WHERE student_id = ? AND date = ? AND (exit_time IS NULL OR exit_time = '')
            ORDER BY visit_id DESC
            LIMIT 1
            """,
            (student_id, target_date),
        ).fetchone()
        if row is None:
            return None
        conn.execute(
            "UPDATE visits SET exit_time = ? WHERE visit_id = ?",
            (now.strftime("%H:%M:%S"), row["visit_id"]),
        )
    return {
        "visit_id": str(row["visit_id"]).zfill(5),
        "student_id": row["student_id"],
        "name": row["name"],
        "father_name": row["father_name"] or "",
        "date": row["date"],
        "entry_time": row["entry_time"],
        "exit_time": now.strftime("%H:%M:%S"),
    }
